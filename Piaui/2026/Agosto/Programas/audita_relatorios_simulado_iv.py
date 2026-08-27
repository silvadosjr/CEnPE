"""Conferência dos resultados e dos apêndices reproduzidos nos relatórios."""
import importlib.util
import json
import hashlib
from pathlib import Path
from zipfile import ZipFile
import numpy as np
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('generator', ROOT / 'Programas/gera_relatorios_simulado_iv.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
n = module.namespace
reference = ROOT.parent / 'Julho/Relatorios/2EM_LP/Relatorio_Tecnico_2EM_LP_S3_TCM_TRI.docx'
with ZipFile(reference) as z:
    inventory = {name: hashlib.sha256(z.read(name)).hexdigest() for name in z.namelist()}
result = {'reference_sha256': hashlib.sha256(reference.read_bytes()).hexdigest(), 'reference_parts': inventory, 'profiles': {}}
for profile in n['PROFILES']:
    data = n['read_profile'](profile)
    m, items, scores, theta, dist = [data[k] for k in ['measures', 'items', 'scores', 'base_theta', 'dist']]
    assert len(scores) == len(theta) == int(dist['N de alunos'].sum())
    assert np.isclose(scores.iloc[:, 1].mean(), m.Dificuldade.sum(), atol=1e-8)
    assert np.isclose(theta.Theta.mean(), data['theta'].loc['media'])
    assert np.isclose(theta.Theta.std(), data['theta'].loc['dp'])
    assert np.allclose(items.bSAEB, 250 + 50 * items.b)
    assert np.allclose(dist.Percentual, dist['N de alunos'] / len(theta) * 100, atol=.00051)
    excluded = sorted(set(m.Item) - set(items.Item))
    assert excluded == ([24] if profile == '2EM_MT' else [])
    assert (2 not in set(m.Item)) == (profile == '3EM_MT')
    file = ROOT / 'Relatorios' / profile / f'Relatorio_Tecnico_{profile}_S4_TCM_TRI.docx'
    doc = Document(file)
    text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells] + [p.text for s in doc.sections for p in s.header.paragraphs])
    for stale in ['Simulado III', 'SIMULADO III', 'Simulado_III', '_S3', 'item 25 anulado', 'item 6 retirado', 'bimodal']:
        assert stale not in text, (profile, stale)
    assert n['fmt_int'](len(scores)) in text
    assert n['fmt_num'](data['theta'].loc['media'], 1) in text
    tcm_table, tri_table, dist_table = doc.tables[-3:]
    assert len(tcm_table.rows) == len(m) + 1
    assert len(tri_table.rows) == len(items) + 1
    for row, (_, src) in zip(tcm_table.rows[1:], m.iterrows()):
        assert [c.text for c in row.cells] == [str(int(src.Item)).zfill(3)] + [n['fmt_num'](src[k], 3) for k in ['Dificuldade','Discriminacao','CBisserial','CPBisserial']]
    for row, (_, src) in zip(tri_table.rows[1:], items.iterrows()):
        expected = [str(int(src.Item)).zfill(3), n['fmt_num'](src.a, 3), n['fmt_num'](src.b, 3), n['fmt_num'](src.bSAEB, 1), n['fmt_num'](src.c, 3), 'Sim' if src.SAEB == 'Sim' else 'Não']
        assert [c.text for c in row.cells] == expected
    for row, (label, src) in zip(dist_table.rows[1:], dist.iterrows()):
        assert [c.text for c in row.cells] == [str(label), n['fmt_int'](src['N de alunos']), n['fmt_num'](src.Percentual, 3) + '%']
    with ZipFile(file) as z:
        for part in ['word/styles.xml','word/theme/theme1.xml','word/numbering.xml','word/footer1.xml']:
            assert hashlib.sha256(z.read(part)).hexdigest() == inventory[part], (profile, part)
    result['profiles'][profile] = {'N':len(scores), 'TCM':len(m), 'TRI':len(items), 'excluded_TRI':excluded, 'mean_score':scores.iloc[:,1].mean(), 'mean_theta':theta.Theta.mean(), 'passed':True}
out = ROOT / 'Relatorios/qa/auditoria.json'
out.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps(result['profiles'], indent=2, ensure_ascii=False))
