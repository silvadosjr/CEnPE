from __future__ import annotations

import os
import tempfile
from datetime import date
from pathlib import Path

import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROFILES = ["2EM_LP", "2EM_MT", "3EM_LP", "3EM_MT"]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "5F6B76"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9DEE5"
CAUTION = "FFF4CE"
RISK = "FDE9E7"

PROFILE_NOTES = {
    "2EM_LP": {
        "lead": (
            "Leitura principal. O Simulado III de Língua Portuguesa da 2ª série apresentou "
            "um item retirado da calibração TRI: o item 5, com possível conflito de gabarito. "
            "O item 13 também apareceu como muito difícil e com discriminação abaixo de 0,20, "
            "mas foi mantido na TRI; sua interpretação deve permanecer acompanhada."
        ),
        "diagnostic_items": [5, 13],
        "status": {
            5: "Excluído da TRI; possível conflito de gabarito",
            13: "Mantido na TRI; dificuldade elevada e discriminação baixa",
        },
        "narrative": (
            "O item 5 concentra a ressalva técnica da aplicação: o gabarito E teve apenas 13,4% "
            "das escolhas e ponto-bisserial baixo, enquanto alternativas não oficiais, sobretudo B "
            "e D, concentraram volume relevante de respostas. O item 13 também foi muito difícil "
            "e pouco discriminativo, mas preservou correlação ponto-bisserial acima de 0,20 e foi "
            "mantido na estimação."
        ),
        "recommendations": [
            "Revisar formalmente o gabarito e o enunciado do item 5 antes de qualquer reutilização.",
            "Manter o item 13 sob acompanhamento pedagógico e psicométrico, por dificuldade extrema e discriminação baixa.",
            "Combinar a leitura TCM com os ajustes TRI e curvas dos itens antes da decisão final de banco.",
        ],
    },
    "2EM_MT": {
        "lead": (
            "Leitura principal. O Simulado III de Matemática da 2ª série manteve todos os 26 itens "
            "na calibração TRI. A TCM apontou itens fracos, especialmente 5 e 2, mas sem retirada; "
            "o item 5 merece atenção por combinar baixa proporção de acertos, baixa discriminação "
            "e uma alternativa não oficial mais atraente."
        ),
        "diagnostic_items": [2, 5],
        "status": {
            2: "Mantido na TRI; ponto-bisserial abaixo de 0,20",
            5: "Mantido na TRI; baixa discriminação e distrator concorrente",
        },
        "narrative": (
            "Nenhum item foi excluído da TRI. Mesmo assim, o item 5 exige revisão: o gabarito E "
            "teve 19,6% das escolhas e ponto-bisserial 0,145, enquanto o distrator D reuniu 31,8% "
            "e ponto-bisserial 0,247. O item 2 apresentou ponto-bisserial abaixo de 0,20, mas "
            "a alternativa correta ainda foi a mais escolhida."
        ),
        "recommendations": [
            "Manter os 26 itens na base analítica, registrando ressalva para os itens 2 e 5.",
            "Revisar pedagogicamente o item 5, com atenção ao distrator D.",
            "Acompanhar os itens com bSAEB acima de 400 e discriminação TRI baixa antes de reaplicações.",
        ],
    },
    "3EM_LP": {
        "lead": (
            "Leitura principal. O Simulado III de Língua Portuguesa da 3ª série retirou o item 5 "
            "da calibração TRI por possível conflito de gabarito e instabilidade no processo de estimação. "
            "Os itens 7 e 8 também apresentaram sinais de funcionamento fraco, mas foram mantidos."
        ),
        "diagnostic_items": [5, 7, 8],
        "status": {
            5: "Excluído da TRI; possível conflito de gabarito e instabilidade",
            7: "Mantido na TRI; discriminação e ponto-bisserial baixos",
            8: "Mantido na TRI; possível conflito de gabarito",
        },
        "narrative": (
            "O item 5 apresentou o sinal mais claro de conflito: o gabarito A reuniu 26,0% das escolhas "
            "e ponto-bisserial 0,095, enquanto o distrator B reuniu 42,0% e ponto-bisserial 0,305. "
            "O item 8 seguiu padrão semelhante, com gabarito E em 15,3% e distrator C em 38,0%, "
            "também com maior associação ao desempenho. O item 7 foi mantido com ressalva por baixa "
            "discriminação e baixa associação ponto-bisserial."
        ),
        "recommendations": [
            "Revisar formalmente o gabarito e o conteúdo dos itens 5 e 8.",
            "Registrar ressalva técnica para os itens 7 e 8, que permaneceram na TRI apesar dos sinais TCM.",
            "Analisar as curvas características e os índices de ajuste antes de aprovar esses itens para banco.",
        ],
    },
    "3EM_MT": {
        "lead": (
            "Leitura principal. O Simulado III de Matemática da 3ª série teve o item 25 anulado "
            "antes da leitura TCM/TRI e o item 6 retirado da calibração TRI por possível conflito de gabarito. "
            "O restante da prova foi calibrado com 24 itens."
        ),
        "diagnostic_items": [6],
        "status": {
            6: "Excluído da TRI; possível conflito de gabarito",
            25: "Anulado antes da TCM/TRI",
        },
        "narrative": (
            "O item 6 apresentou funcionamento incompatível com a chave: o gabarito A reuniu 12,2% "
            "das escolhas e ponto-bisserial praticamente nulo, enquanto o distrator D concentrou 51,9% "
            "das respostas e ponto-bisserial 0,281. O padrão sustenta a retirada do item da calibração TRI. "
            "O item 25 foi tratado como anulado, reduzindo o conjunto TCM a 25 itens."
        ),
        "recommendations": [
            "Manter documentada a anulação do item 25 e a exclusão TRI do item 6.",
            "Revisar formalmente o gabarito, a resolução e os distratores do item 6 antes de qualquer uso posterior.",
            "Interpretar os indicadores da aplicação com denominador de 25 itens na TCM e 24 na TRI.",
        ],
    },
}


def fmt_num(value: float, decimals: int = 1) -> str:
    return f"{float(value):,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_int(value: int) -> str:
    return f"{int(value):,}".replace(",", ".")


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_shading(paragraph, fill, border=BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_source(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"Fonte: {text}")
    set_font(run, size=8.5, color=MUTED)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_font(run, size=9, color=MUTED, italic=True)


def add_picture_with_alt(doc, path, width, alt_text):
    shape = doc.add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)


def add_table(doc, headers, rows, widths_dxa, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(label))
        set_font(run, size=9, color=NAVY, bold=True)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_font(run, size=9.2, color="222222")
    set_table_geometry(table, widths_dxa)
    return table


def configure_document(doc, profile):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        st = styles[style_name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        st = styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run(f"PIAUÍ 2026 | SIMULADO III | {profile.replace('_', '-')}")
    set_font(run, size=8.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])


def read_profile(profile):
    series = int(profile[0])
    discipline = profile.split("_")[1]
    tcm = ROOT / f"ResultadosTCM_{profile}" / "Simulado_III"
    tri = ROOT / f"ResultadosTRI_{profile}" / "Simulado_III"
    data = {
        "scores": pd.read_excel(tcm / "EscoresBrutos_Aluno.xlsx"),
        "measures": pd.read_excel(tcm / "MedidasTCM_Item.xlsx"),
        "theta": pd.read_excel(tri / f"ResumoTheta_{profile}_S3.xlsx", index_col=0).iloc[:, 0],
        "theta_ic": pd.read_excel(tri / f"ResumoThetaIC_{profile}_S3.xlsx", index_col=0),
        "items": pd.read_excel(tri / f"EstItens_{profile}_S3.xlsx"),
        "dist": pd.read_excel(tri / f"DistriAlunosClasse{discipline}{series}serie_S3.xlsx", index_col=0),
        "choice": pd.read_excel(tcm / "mDificNRDF.xlsx", index_col=0),
        "disc_alt": pd.read_excel(tcm / "mDiscNRDF.xlsx", index_col=0),
        "pbis_alt": pd.read_excel(tcm / "mcpBisNRDF.xlsx", index_col=0),
    }
    return data


def make_charts(profile, data, workdir):
    try:
        font = ImageFont.truetype("arial.ttf", 20)
        small = ImageFont.truetype("arial.ttf", 14)
        tiny = ImageFont.truetype("arial.ttf", 11)
    except OSError:
        font = small = tiny = ImageFont.load_default()

    measures = data["measures"]
    problem_indices = set(PROFILE_NOTES[profile]["diagnostic_items"])
    color = "#2E74B5"

    def panel(draw, box, title, values, vmax, reference=None, shade=None):
        x0, y0, x1, y1 = box
        draw.text((x0, y0), title, font=font, fill="#203748")
        top, bottom, left, right = y0 + 36, y1 - 34, x0 + 48, x1 - 12
        if shade:
            lo, hi = shade
            sy1 = bottom - (hi / vmax) * (bottom - top)
            sy0 = bottom - (lo / vmax) * (bottom - top)
            draw.rectangle((left, sy1, right, sy0), fill="#E8EEF5")
        for tick in range(0, 5):
            value = vmax * tick / 4
            yy = bottom - (value / vmax) * (bottom - top)
            draw.line((left, yy, right, yy), fill="#D9DEE5", width=1)
            draw.text((x0, yy - 7), f"{value:.2f}".replace(".", ","), font=tiny, fill="#5F6B76")
        if reference is not None:
            yy = bottom - (reference / vmax) * (bottom - top)
            for xx in range(left, right, 10):
                draw.line((xx, yy, min(xx + 5, right), yy), fill="#607D8B", width=2)
        draw.line((left, top, left, bottom), fill="#5F6B76", width=1)
        draw.line((left, bottom, right, bottom), fill="#5F6B76", width=1)
        gap = (right - left) / len(values)
        bw = max(4, int(gap * 0.68))
        for idx, value in enumerate(values, start=1):
            cx = left + (idx - 0.5) * gap
            fill = "#B3261E" if idx in problem_indices else color
            value = float(value)
            yy = bottom - (max(value, 0.0) / vmax) * (bottom - top)
            if value >= 0:
                draw.rectangle((cx - bw / 2, yy, cx + bw / 2, bottom), fill=fill)
            else:
                draw.rectangle((cx - bw / 2, bottom, cx + bw / 2, bottom + 12), fill=fill)
            if idx % 2 == 0:
                draw.text((cx - 7, bottom + 5), str(idx), font=tiny, fill="#5F6B76")

    img = Image.new("RGB", (1800, 700), "white")
    draw = ImageDraw.Draw(img)
    panel(draw, (40, 30, 880, 660), "Proporção de acertos por item", measures["Dificuldade"].tolist(), 0.90, shade=(0.2, 0.4))
    panel(draw, (920, 30, 1760, 660), "Discriminação por item", measures["Discriminacao"].tolist(), 0.75, reference=0.20)
    tcm_chart = workdir / f"{profile}_tcm_itens.png"
    img.save(tcm_chart, dpi=(180, 180))

    dist = data["dist"]
    classes = [str(x) for x in dist.index if float(dist.loc[x, "N de alunos"]) > 0]
    labels = []
    for c in classes:
        labels.append("350+" if c == "[350,1e+03)" else c.replace("[", "").replace(")", "").replace(",", "–"))
    values = [float(dist.loc[c, "Percentual"]) for c in classes]
    img = Image.new("RGB", (1800, 760), "white")
    draw = ImageDraw.Draw(img)
    draw.text((60, 30), "Distribuição dos estudantes por faixa de proficiência", font=font, fill="#203748")
    left, right, top, bottom = 120, 1740, 100, 650
    vmax = max(30, int(max(values) / 5 + 2) * 5)
    for tick in range(0, vmax + 1, 5):
        yy = bottom - (tick / vmax) * (bottom - top)
        draw.line((left, yy, right, yy), fill="#D9DEE5", width=1)
        draw.text((65, yy - 8), str(tick), font=small, fill="#5F6B76")
    draw.line((left, top, left, bottom), fill="#5F6B76", width=2)
    draw.line((left, bottom, right, bottom), fill="#5F6B76", width=2)
    group = (right - left) / len(labels)
    bw = int(group * 0.55)
    for idx, label in enumerate(labels):
        cx = left + (idx + 0.5) * group
        value = values[idx]
        yy = bottom - (value / vmax) * (bottom - top)
        draw.rectangle((cx - bw / 2, yy, cx + bw / 2, bottom), fill="#2E74B5")
        draw.text((cx - 35, bottom + 12), label, font=tiny, fill="#5F6B76")
    theta_chart = workdir / f"{profile}_theta_dist.png"
    img.save(theta_chart, dpi=(180, 180))
    return tcm_chart, theta_chart


def discipline_name(profile):
    return "Matemática" if profile.endswith("MT") else "Língua Portuguesa"


def series_ordinal(profile):
    return "2ª" if profile.startswith("2") else "3ª"


def build_report(profile, data, charts):
    doc = Document()
    configure_document(doc, profile)
    tcm_chart, theta_chart = charts
    notes = PROFILE_NOTES[profile]
    series = int(profile[0])
    discipline = profile.split("_")[1]

    # Cover
    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RELATÓRIO TÉCNICO")
    set_font(run, size=11, color=BLUE, bold=True)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Análises TCM e TRI")
    set_font(run, size=30, color=NAVY, bold=True)
    p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{series_ordinal(profile)} série do Ensino Médio – {discipline_name(profile)}")
    set_font(run, size=15, color=DARK_BLUE)
    p.paragraph_format.space_after = Pt(28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Simulado III | Piauí, 2026")
    set_font(run, size=11, color=MUTED, bold=True)
    p.paragraph_format.space_after = Pt(34)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documento-base metodológico: Relatório 3EM-MT 2025 e relatórios refinados de Junho/2026")
    set_font(run, size=10.5, color=MUTED, italic=True)
    p.paragraph_format.space_after = Pt(42)
    months = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
    today = date.today()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{today.day} de {months[today.month - 1]} de {today.year}")
    set_font(run, size=10, color=MUTED)
    doc.add_page_break()

    scores = data["scores"]
    measures = data["measures"]
    theta = data["theta"]
    theta_ic = data["theta_ic"].loc["SAEB"]
    items = data["items"]
    dist = data["dist"]
    mean_score = float(scores.iloc[:, 1].mean())
    tcm_count = len(measures)
    tri_count = len(items)
    tcm_items = set(measures["Item"].astype(int))
    tri_items = set(items["Item"].astype(int))
    excluded = sorted(tcm_items - tri_items)
    exclusion_text = "Nenhum" if not excluded else ", ".join(str(x) for x in excluded)
    if profile == "3EM_MT":
        exclusion_text = "6; item 25 anulado"

    doc.add_heading("1. Síntese executiva", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    set_paragraph_shading(p, "EDF4FB", BLUE)
    run = p.add_run(notes["lead"])
    set_font(run, size=10.5, color=NAVY, bold=True)

    summary_rows = [[
        "Simulado III",
        fmt_int(len(scores)),
        f"{tcm_count} / {tri_count}",
        f"{fmt_num(mean_score, 2)} / {tcm_count}",
        f"{fmt_num(mean_score / tcm_count * 100, 1)}%",
        fmt_num(theta.loc["media"], 1),
        exclusion_text,
    ]]
    add_table(
        doc,
        ["Aplicação", "N", "Itens TCM/TRI", "Escore médio", "% do total", "Proficiência média", "Exclusão TRI"],
        summary_rows,
        [1150, 900, 1350, 1250, 1000, 1700, 2010],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_source(doc, "EscoresBrutos_Aluno, MedidasTCM_Item, EstItens e ResumoTheta do Simulado III.")

    doc.add_heading("2. Dados e método", level=1)
    doc.add_paragraph(
        f"A análise utiliza os arquivos de resultados do Simulado III de {discipline_name(profile)} da {series_ordinal(profile)} série do Ensino Médio, "
        "armazenados na pasta Julho. A TCM foi usada para examinar proporção de acertos, discriminação e correlações bisserial e ponto-bisserial. "
        "A TRI foi estimada pelo modelo logístico de três parâmetros, com equalização para a escala SAEB quando indicado nos arquivos de itens."
    )
    doc.add_paragraph(
        "Como nos relatórios anteriores, os limiares de discriminação e correlação ponto-bisserial abaixo de 0,20 são tratados como sinais diagnósticos, "
        "não como regras automáticas de exclusão. A decisão final combina TCM, parâmetros TRI, ajustes do modelo, estabilidade da estimação e revisão pedagógica."
    )

    doc.add_heading("3. Resultados da TCM", level=1)
    add_picture_with_alt(
        doc,
        tcm_chart,
        Inches(6.35),
        "Gráficos de barras mostram proporção de acertos e discriminação por item no Simulado III; itens com ressalva aparecem em vermelho.",
    )
    add_caption(doc, "Figura 1. Proporção de acertos e discriminação por item.")
    add_source(doc, "MedidasTCM_Item.xlsx, Simulado III.")

    weak = measures[(measures["Discriminacao"] < 0.20) | (measures["CPBisserial"] < 0.20)].copy()
    issue_items = sorted(set(notes["diagnostic_items"]).union(weak["Item"].astype(int).tolist()))
    issue_rows = []
    for item_number in issue_items:
        measure_match = measures[measures["Item"].astype(int) == item_number]
        if measure_match.empty:
            issue_rows.append([str(item_number).zfill(3), "Anulado", "—", "—", "—", notes["status"].get(item_number, "Anulado")])
            continue
        row = measure_match.iloc[0]
        issue_rows.append([
            str(item_number).zfill(3),
            fmt_num(row["Dificuldade"], 3),
            fmt_num(row["Discriminacao"], 3),
            fmt_num(row["CBisserial"], 3),
            fmt_num(row["CPBisserial"], 3),
            notes["status"].get(item_number, "Ressalva técnica"),
        ])
    add_table(
        doc,
        ["Item", "Prop. acertos", "Discriminação", "Corr. bisserial", "Corr. ponto-bisserial", "Encaminhamento"],
        issue_rows,
        [750, 1350, 1350, 1500, 1700, 2710],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5 + [WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_source(doc, "MedidasTCM_Item.xlsx e Impressões gerais..odt.")
    doc.add_paragraph(notes["narrative"])

    doc.add_heading("Diagnóstico das alternativas", level=2)
    alt_rows = []
    for item_number in issue_items:
        if item_number not in data["choice"].index and item_number > len(data["choice"]):
            continue
        if item_number == 25 and profile == "3EM_MT":
            continue
        row_choice = data["choice"].iloc[item_number - 1]
        row_disc = data["disc_alt"].iloc[item_number - 1]
        row_pbis = data["pbis_alt"].iloc[item_number - 1]
        for alternative in ["A", "B", "C", "D", "E"]:
            alt_rows.append([
                str(item_number),
                alternative,
                fmt_num(row_choice[alternative], 1) + "%",
                fmt_num(row_disc[alternative], 3),
                fmt_num(row_pbis[alternative], 3),
                "Gabarito" if alternative == row_choice["Gabarito"] else "Distrator",
            ])
    if alt_rows:
        add_table(
            doc,
            ["Item", "Alternativa", "% escolha", "Discriminação", "Ponto-bisserial", "Papel"],
            alt_rows,
            [800, 1150, 1400, 1700, 1900, 2410],
            [WD_ALIGN_PARAGRAPH.CENTER] * 6,
        )
        add_source(doc, "mDificNRDF, mDiscNRDF e mcpBisNRDF, Simulado III.")

    doc.add_heading("4. Resultados da TRI", level=1)
    tri_rows = [[
        "Simulado III",
        fmt_num(theta.loc["media"], 1),
        f"{fmt_num(theta_ic['IC_Inf'], 1)} a {fmt_num(theta_ic['IC_Sup'], 1)}",
        fmt_num(theta.loc["dp"], 1),
        fmt_num(theta.loc["med."], 1),
        f"{fmt_num(theta.loc['1o Q'], 1)} a {fmt_num(theta.loc['3oQ'], 1)}",
        str(int((items["SAEB"] == "Sim").sum())),
    ]]
    add_table(
        doc,
        ["Aplicação", "Média", "IC 95% da média", "DP EAP", "Mediana", "Intervalo interquartil", "Itens fixados"],
        tri_rows,
        [1300, 950, 1700, 1000, 1050, 1900, 1460],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_source(doc, "ResumoTheta e ResumoThetaIC.")

    add_picture_with_alt(
        doc,
        theta_chart,
        Inches(6.35),
        "Gráfico de barras mostra os percentuais de estudantes por faixa de proficiência no Simulado III.",
    )
    add_caption(doc, "Figura 2. Distribuição percentual dos estudantes por faixa de proficiência.")
    add_source(doc, f"DistriAlunosClasse{discipline}{series}serie_S3.xlsx.")

    doc.add_paragraph(
        f"A proficiência média estimada foi {fmt_num(theta.loc['media'], 1)}, com mediana {fmt_num(theta.loc['med.'], 1)} "
        f"e desvio-padrão EAP de {fmt_num(theta.loc['dp'], 1)}. A proporção de estudantes na faixa de 350 pontos ou mais foi "
        f"{fmt_num(dist.iloc[-1]['Percentual'], 1)}%."
    )

    doc.add_heading("Parâmetros dos itens", level=2)
    param_rows = []
    for signal, subset in [
        ("bSAEB > 400", items[items["bSAEB"] > 400]),
        ("a < 0,60", items[items["a"] < 0.60]),
        ("a > 4,00", items[items["a"] > 4.00]),
        ("c > 0,25", items[items["c"] > 0.25]),
    ]:
        if not subset.empty:
            param_rows.append([signal, ", ".join(str(int(x)) for x in subset["Item"].tolist())])
    if param_rows:
        add_table(
            doc,
            ["Sinal para revisão", "Itens"],
            param_rows,
            [2600, 6760],
            [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        )
        add_source(doc, f"EstItens_{profile}_S3.xlsx. Os pontos de corte são diagnósticos e não implicam exclusão automática.")
    else:
        doc.add_paragraph("Não foram observados itens nos pontos de corte diagnósticos selecionados para parâmetros TRI extremos.")

    doc.add_heading("5. Recomendações", level=1)
    for recommendation in notes["recommendations"]:
        doc.add_paragraph(recommendation, style="List Number")

    doc.add_heading("6. Referências e arquivos-fonte", level=1)
    references = [
        "Primi, R. (2012). Psicometria: fundamentos matemáticos da Teoria Clássica dos Testes. Avaliação Psicológica, 11, 297–307.",
        "Wu, M.; Tam, H. P.; Jen, T.-H. (2017). Educational Measurement for Applied Researchers: Theory into Practice. Springer.",
        "Andrade, D. F.; Tavares, H. R.; Valle, R. C. (2000). Teoria da Resposta ao Item: conceitos e aplicações. ABE.",
        "Baker, F. B.; Kim, S.-H. (2004). Item Response Theory: Parameter Estimation Techniques. 2. ed. Marcel Dekker.",
        "Relatório Técnico 3EM-MT 2025 e relatórios refinados de Junho/2026.",
    ]
    for ref in references:
        doc.add_paragraph(ref)
    doc.add_heading("Arquivos de resultados utilizados", level=2)
    for src in [
        f"ResultadosTCM_{profile}/Simulado_III: EscoresBrutos_Aluno.xlsx, MedidasTCM_Item.xlsx, mDificNRDF.xlsx, mDiscNRDF.xlsx e mcpBisNRDF.xlsx.",
        f"ResultadosTRI_{profile}/Simulado_III: ResumoTheta, ResumoThetaIC, EstItens, DistriAlunosClasse e BaseRespTheta.",
        f"Programas/TCM_{profile}_S3.R e Fit_3PL_Fix_{profile}_S3.R.",
        "Impressões gerais..odt, com as ressalvas técnicas registradas pela equipe.",
    ]:
        doc.add_paragraph(src, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Apêndice A. Medidas TCM por item", level=1)
    tcm_rows = []
    for _, row in measures.iterrows():
        tcm_rows.append([
            str(int(row["Item"])).zfill(3),
            fmt_num(row["Dificuldade"], 3),
            fmt_num(row["Discriminacao"], 3),
            fmt_num(row["CBisserial"], 3),
            fmt_num(row["CPBisserial"], 3),
        ])
    add_table(
        doc,
        ["Item", "Proporção de acertos", "Discriminação", "Corr. bisserial", "Corr. ponto-bisserial"],
        tcm_rows,
        [900, 2200, 1800, 2000, 2460],
        [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    add_source(doc, f"ResultadosTCM_{profile}/Simulado_III/MedidasTCM_Item.xlsx.")

    doc.add_page_break()
    doc.add_heading("Apêndice B. Parâmetros TRI por item", level=1)
    tri_item_rows = []
    for _, row in items.iterrows():
        tri_item_rows.append([
            str(int(row["Item"])).zfill(3),
            fmt_num(row["a"], 3),
            fmt_num(row["b"], 3),
            fmt_num(row["bSAEB"], 1),
            fmt_num(row["c"], 3),
            "Sim" if row["SAEB"] == "Sim" else "Não",
        ])
    add_table(
        doc,
        ["Item", "a", "b", "b (escala 250,50)", "c", "Fixado"],
        tri_item_rows,
        [900, 1300, 1300, 2500, 1400, 1960],
        [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_source(doc, f"ResultadosTRI_{profile}/Simulado_III/EstItens_{profile}_S3.xlsx.")

    doc.add_page_break()
    doc.add_heading("Apêndice C. Distribuição por faixa de proficiência", level=1)
    dist_rows = []
    for classe, row in dist.iterrows():
        dist_rows.append([str(classe), fmt_int(row["N de alunos"]), fmt_num(row["Percentual"], 3) + "%"])
    add_table(
        doc,
        ["Faixa", "N de alunos", "Percentual"],
        dist_rows,
        [4200, 2200, 2960],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_source(doc, f"DistriAlunosClasse{discipline}{series}serie_S3.xlsx.")

    out_dir = ROOT / "Relatorios" / profile
    out_dir.mkdir(parents=True, exist_ok=True)
    output = out_dir / f"Relatorio_Tecnico_{profile}_S3_TCM_TRI.docx"
    doc.save(output)
    return output


def main():
    selected = os.environ.get("REPORT_PROFILE")
    profiles = [selected] if selected else PROFILES
    outputs = []
    with tempfile.TemporaryDirectory(prefix="relatorios_s3_") as tmp:
        tmp_path = Path(tmp)
        for profile in profiles:
            data = read_profile(profile)
            charts = make_charts(profile, data, tmp_path)
            outputs.append(build_report(profile, data, charts))
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
